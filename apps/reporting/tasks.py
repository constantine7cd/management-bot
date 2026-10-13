import csv
import io
from datetime import date, timedelta

from celery import shared_task
from django.core.files.base import ContentFile
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.auditlog.models import record
from apps.leave_management.models import LeaveRequest

from .models import ExportJob
from .services import accessible_employees, accessible_shifts


def _range(job):
    end = date.fromisoformat(job.filters.get("end", str(date.today())))
    start = date.fromisoformat(job.filters.get("start", str(end - timedelta(days=30))))
    return start, end


def _shift_rows(job):
    start, end = _range(job)
    shifts = accessible_shifts(job.requested_by).filter(
        started_at__date__gte=start - timedelta(days=1),
        started_at__date__lte=end + timedelta(days=1),
    )
    for shift in shifts.iterator():
        local_start = shift.local_started_at
        if not start <= local_start.date() <= end:
            continue
        local_end = shift.local_ended_at
        yield [
            str(shift.employee),
            local_start.isoformat(),
            shift.start_timezone,
            local_end.isoformat() if local_end else "",
            shift.end_timezone or shift.start_timezone if local_end else "",
            shift.started_at.isoformat(),
            shift.ended_at.isoformat() if shift.ended_at else "",
            round(shift.duration.total_seconds() / 3600, 2) if shift.duration else "",
            shift.start_workplace.name,
            shift.end_workplace.name if shift.end_workplace else "",
            str(shift.workplace_shift or ""),
        ]


def _leave_rows(job):
    start, end = _range(job)
    requests = LeaveRequest.objects.filter(
        employee__in=accessible_employees(job.requested_by),
        start_date__lte=end,
        end_date__gte=start,
    ).select_related("employee", "leave_type")
    for leave in requests.iterator():
        yield [
            str(leave.employee),
            str(leave.leave_type),
            leave.start_date.isoformat(),
            leave.end_date.isoformat(),
            leave.requested_days,
            leave.get_status_display(),
            leave.reason,
        ]


SHIFT_HEADERS = [
    "Employee",
    "Started (local)",
    "Start timezone",
    "Ended (local)",
    "End timezone",
    "Started (UTC)",
    "Ended (UTC)",
    "Hours",
    "Start workplace",
    "End workplace",
    "Daily shift",
]
LEAVE_HEADERS = ["Employee", "Type", "Start date", "End date", "Work days", "Status", "Reason"]


def _format_sheet(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4338CA")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


@shared_task(bind=True, autoretry_for=(OSError,), retry_backoff=True, max_retries=3)
def generate_export(self, job_id):
    job = ExportJob.objects.select_related("requested_by").get(pk=job_id)
    job.status = ExportJob.Status.RUNNING
    job.save(update_fields=["status"])
    try:
        if job.kind == ExportJob.Kind.ATTENDANCE_CSV:
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(["Record type", *SHIFT_HEADERS, *LEAVE_HEADERS[1:]])
            for row in _shift_rows(job):
                writer.writerow(["Attendance", *row, *([""] * (len(LEAVE_HEADERS) - 1))])
            for row in _leave_rows(job):
                writer.writerow(
                    ["Day off", row[0], *([""] * (len(SHIFT_HEADERS) - 1)), *row[1:]]
                )
            content = ContentFile(("\ufeff" + output.getvalue()).encode("utf-8"))
            filename = "attendance.csv"
        elif job.kind == ExportJob.Kind.VACATIONS_CSV:
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(LEAVE_HEADERS)
            writer.writerows(_leave_rows(job))
            content = ContentFile(("\ufeff" + output.getvalue()).encode("utf-8"))
            filename = "vacations-and-days-off.csv"
        else:
            workbook = Workbook()
            sheet = workbook.active
            if job.kind == ExportJob.Kind.EMPLOYEES_XLSX:
                sheet.title = "Employees"
            elif job.kind == ExportJob.Kind.VACATIONS_XLSX:
                sheet.title = "Vacations and days off"
            else:
                sheet.title = "Attendance"
            if job.kind == ExportJob.Kind.EMPLOYEES_XLSX:
                sheet.append(
                    ["Employee", "Category", "Email", "Team", "Rate", "Currency", "Active"]
                )
                for employee in accessible_employees(job.requested_by).iterator():
                    sheet.append(
                        [
                            str(employee),
                            employee.get_category_display(),
                            employee.email,
                            employee.team.name if employee.team else "",
                            employee.hourly_rate,
                            employee.rate_currency,
                            employee.is_active,
                        ]
                    )
                filename = "employees.xlsx"
            elif job.kind == ExportJob.Kind.VACATIONS_XLSX:
                sheet.append(LEAVE_HEADERS)
                for row in _leave_rows(job):
                    sheet.append(row)
                filename = "vacations-and-days-off.xlsx"
            else:
                sheet.append(SHIFT_HEADERS)
                for row in _shift_rows(job):
                    sheet.append(row)
                leave_sheet = workbook.create_sheet("Days off")
                leave_sheet.append(LEAVE_HEADERS)
                for row in _leave_rows(job):
                    leave_sheet.append(row)
                _format_sheet(leave_sheet)
                filename = "attendance.xlsx"
            _format_sheet(sheet)
            binary = io.BytesIO()
            workbook.save(binary)
            content = ContentFile(binary.getvalue())
        job.file.save(filename, content, save=False)
        job.status = ExportJob.Status.READY
        job.save(update_fields=["file", "status"])
        record("reporting.export_ready", job, actor=job.requested_by, metadata={"kind": job.kind})
    except Exception as error:
        job.status = ExportJob.Status.FAILED
        job.error = str(error)[:1000]
        job.save(update_fields=["status", "error"])
        raise


@shared_task
def expire_exports():
    jobs = ExportJob.objects.filter(expires_at__lt=timezone.now()).exclude(status=ExportJob.Status.EXPIRED)
    count = 0
    for job in jobs:
        if job.file:
            job.file.delete(save=False)
        job.status = ExportJob.Status.EXPIRED
        job.save(update_fields=["file", "status"])
        count += 1
    return count
