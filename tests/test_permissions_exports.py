import csv
from datetime import date, datetime
from io import StringIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from apps.attendance.models import Shift
from apps.leave_management.models import LeaveRequest, LeaveType
from apps.organization.models import Employee, Holiday, Team, Workplace
from apps.reporting.models import ExportJob
from apps.reporting.tasks import generate_export

pytestmark = pytest.mark.django_db


def test_manager_has_company_wide_employee_access(client, manager, employee):
    other_team = Team.objects.create(name="Finance")
    other = Employee.objects.create(
        first_name="Ion", last_name="Ionescu", team=other_team, hourly_rate="999.00"
    )
    client.force_login(manager)
    response = client.get("/employees/")
    assert response.status_code == 200
    assert b"125.50 MDL" in response.content
    assert b"999.00 MDL" in response.content
    assert client.get(f"/employees/{employee.id}/edit/").status_code == 200
    assert client.get(f"/employees/{other.id}/edit/").status_code == 200


def test_manager_employee_export_contains_all_company_employees(
    manager, employee, tmp_path, settings
):
    settings.MEDIA_ROOT = tmp_path
    other_team = Team.objects.create(name="Finance")
    Employee.objects.create(first_name="Ion", last_name="Secret", team=other_team, hourly_rate="999.00")
    job = ExportJob.objects.create(requested_by=manager, kind=ExportJob.Kind.EMPLOYEES_XLSX, filters={"start": str(date.today()), "end": str(date.today())})
    generate_export(str(job.id))
    job.refresh_from_db()
    workbook = load_workbook(job.file.path)
    values = list(workbook.active.values)
    assert values[0][1] == "Category"
    assert any(row[1] == "Operations worker" for row in values[1:])
    assert any("Ana Popescu" in row for row in values)
    assert any("Ion Secret" in row for row in values)
    assert job.status == ExportJob.Status.READY


def test_manager_is_blocked_from_admin_only_areas(client, manager):
    client.force_login(manager)
    exports_response = client.get("/exports/")
    assert exports_response.status_code == 200
    assert b"Vacations and days off CSV" in exports_response.content
    assert b"Vacations and days off XLSX" in exports_response.content
    assert client.get("/managers/").status_code == 403
    assert client.get("/audit/").status_code == 403
    assert client.get("/django-admin/").status_code == 403


def test_health_endpoints(client):
    assert client.get("/health/live").status_code == 200
    assert client.get("/health/ready").status_code == 200


def test_anonymous_panel_redirects_to_login(client):
    response = client.get("/employees/")
    assert response.status_code == 302
    assert "/accounts/login/" in response.url


def test_non_staff_account_is_denied_panel(client, django_user_model):
    user = django_user_model.objects.create_user(email="outside@example.com", password="test-pass")
    client.force_login(user)
    assert client.get("/employees/").status_code == 403


def test_webhook_rejects_missing_secret(client, settings):
    settings.TELEGRAM_WEBHOOK_SECRET = "secret"
    assert client.post("/telegram/webhook/", data="{}", content_type="application/json").status_code == 403


def test_admin_can_create_manager_and_assign_team(client, admin_user, team):
    client.force_login(admin_user)
    response = client.post(
        "/managers/add/",
        {
            "email": "new.manager@example.com",
            "first_name": "New",
            "last_name": "Manager",
            "teams": [team.id],
            "password": "Strong-initial-password-928!",
            "is_active": "on",
        },
    )
    assert response.status_code == 302
    manager = admin_user.__class__.objects.get(email="new.manager@example.com")
    assert manager.is_staff
    assert manager.managed_teams.get() == team


def test_admin_can_create_workplace_with_numbered_shifts(client, admin_user):
    client.force_login(admin_user)
    response = client.post(
        "/workplaces/add/",
        {
            "name": "Tirnova",
            "latitude": "47.844000",
            "longitude": "28.178000",
            "radius_meters": 150,
            "max_accuracy_meters": 100,
            "is_active": "on",
            "shifts-TOTAL_FORMS": "2",
            "shifts-INITIAL_FORMS": "0",
            "shifts-MIN_NUM_FORMS": "0",
            "shifts-MAX_NUM_FORMS": "1000",
            "shifts-0-number": "1",
            "shifts-0-name": "Morning",
            "shifts-0-start_time": "08:00",
            "shifts-0-end_time": "16:00",
            "shifts-0-is_active": "on",
            "shifts-1-number": "2",
            "shifts-1-name": "Evening",
            "shifts-1-start_time": "16:00",
            "shifts-1-end_time": "00:00",
            "shifts-1-is_active": "on",
        },
    )
    assert response.status_code == 302
    workplace = Workplace.objects.get(name="Tirnova")
    assert list(workplace.daily_shifts.values_list("number", flat=True)) == [1, 2]
    assert workplace.timezone == "Europe/Chisinau"


def test_manager_can_create_employees_and_workplaces(client, manager, team):
    client.force_login(manager)
    team_response = client.post(
        "/teams/add/", {"name": "Manager-created team", "is_active": "on"}
    )
    assert team_response.status_code == 302
    assert Team.objects.filter(name="Manager-created team").exists()

    employee_response = client.post(
        "/employees/add/",
        {
            "first_name": "Maria",
            "last_name": "Rusu",
            "email": "maria@example.com",
            "team": team.id,
            "language": "en",
            "rate_currency": "MDL",
            "holiday_country": "MD",
            "is_active": "on",
        },
    )
    assert employee_response.status_code == 302
    assert Employee.objects.filter(email="maria@example.com").exists()

    workplace_response = client.post(
        "/workplaces/add/",
        {
            "name": "Manager-created office",
            "latitude": "47.010500",
            "longitude": "28.863800",
            "radius_meters": 150,
            "max_accuracy_meters": 100,
            "is_active": "on",
            "shifts-TOTAL_FORMS": "1",
            "shifts-INITIAL_FORMS": "0",
            "shifts-MIN_NUM_FORMS": "0",
            "shifts-MAX_NUM_FORMS": "1000",
            "shifts-0-number": "1",
            "shifts-0-name": "Morning",
            "shifts-0-start_time": "08:00",
            "shifts-0-end_time": "16:00",
            "shifts-0-is_active": "on",
        },
    )
    assert workplace_response.status_code == 302
    assert Workplace.objects.filter(name="Manager-created office").exists()


def test_admin_can_create_team(client, admin_user, manager):
    client.force_login(admin_user)
    response = client.post(
        "/teams/add/",
        {"name": "Night crew", "managers": [manager.id], "is_active": "on"},
    )
    assert response.status_code == 302
    assert Team.objects.get(name="Night crew").managers.get() == manager


def test_admin_can_import_country_holidays(client, admin_user):
    client.force_login(admin_user)
    response = client.post("/holidays/", {"country_code": "MD", "year": 2026})
    assert response.status_code == 302
    assert Holiday.objects.filter(country_code="MD", date__year=2026).exists()


def test_attendance_exports_include_local_time_and_days_off(
    manager, employee, workplace, tmp_path, settings
):
    settings.MEDIA_ROOT = tmp_path
    Shift.objects.create(
        employee=employee,
        started_at=datetime(2026, 10, 12, 6, tzinfo=ZoneInfo("UTC")),
        ended_at=datetime(2026, 10, 12, 14, tzinfo=ZoneInfo("UTC")),
        start_workplace=workplace,
        end_workplace=workplace,
        start_timezone=workplace.timezone,
        end_timezone=workplace.timezone,
    )
    leave_type = LeaveType.objects.create(name_en="Vacation", name_ru="Отпуск")
    LeaveRequest.objects.create(
        employee=employee,
        leave_type=leave_type,
        start_date=date(2026, 10, 13),
        end_date=date(2026, 10, 14),
        requested_days=2,
    )
    filters = {"start": "2026-10-01", "end": "2026-10-31"}
    csv_job = ExportJob.objects.create(
        requested_by=manager, kind=ExportJob.Kind.ATTENDANCE_CSV, filters=filters
    )
    generate_export(str(csv_job.id))
    csv_job.refresh_from_db()
    rows = list(csv.reader(StringIO(csv_job.file.read().decode("utf-8-sig"))))
    assert any(row[0] == "Attendance" and "Europe/Chisinau" in row for row in rows)
    assert any(row[0] == "Day off" and "Vacation" in row for row in rows)

    xlsx_job = ExportJob.objects.create(
        requested_by=manager, kind=ExportJob.Kind.ATTENDANCE_XLSX, filters=filters
    )
    generate_export(str(xlsx_job.id))
    xlsx_job.refresh_from_db()
    workbook = load_workbook(xlsx_job.file.path)
    assert workbook.sheetnames == ["Attendance", "Days off"]
    assert list(workbook["Days off"].values)[1][0] == "Ana Popescu"

    vacation_csv_job = ExportJob.objects.create(
        requested_by=manager, kind=ExportJob.Kind.VACATIONS_CSV, filters=filters
    )
    generate_export(str(vacation_csv_job.id))
    vacation_csv_job.refresh_from_db()
    vacation_rows = list(
        csv.reader(StringIO(vacation_csv_job.file.read().decode("utf-8-sig")))
    )
    assert vacation_rows[0] == [
        "Employee",
        "Type",
        "Start date",
        "End date",
        "Work days",
        "Status",
        "Reason",
    ]
    assert any("Vacation" in row for row in vacation_rows)

    vacation_xlsx_job = ExportJob.objects.create(
        requested_by=manager, kind=ExportJob.Kind.VACATIONS_XLSX, filters=filters
    )
    generate_export(str(vacation_xlsx_job.id))
    vacation_xlsx_job.refresh_from_db()
    vacation_workbook = load_workbook(vacation_xlsx_job.file.path)
    assert vacation_workbook.sheetnames == ["Vacations and days off"]
    assert list(vacation_workbook.active.values)[1][0] == "Ana Popescu"
